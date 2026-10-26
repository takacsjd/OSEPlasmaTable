"""Create the initial human-editable AI_CAD workflow workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = ROOT / "AI_CAD" / "workflow.xlsx"

DEVICE_BRIEF = [
    ["project_id", "OSE-PLASMA-TABLE", "draft", "REF-README", "Stable project identifier."],
    ["device_name", "Uitvindersgilde plasma table", "draft", "REF-README", "Project name; machine design is still evolving."],
    ["purpose", "CNC plasma cutting with a water table; future manual tool changes may add drilling, scribing, and light machining.", "draft", "REF-README", "Keep plasma cutting as the first design priority."],
    ["narrative", "Design a CNC plasma table for a community workshop. It should support at least 1000 x 2000 mm sheet stock within an outer envelope target of X 1600 x Y 2500 mm. The project includes a water table and adjustable table height. Future capabilities under consideration include drilling, scribing, light milling, and a tube-plasma rotary axis.", "draft", "REF-README", "Narrative summarizes current notes; it does not resolve open design decisions."],
    ["in_scope", "Water table, slats, supporting structure, table lift, XY gantry and drives, torch Z/breakaway, control integration, and staged future tools.", "draft", "REF-README", "Treat candidate/future components as proposals, not approved scope."],
    ["out_of_scope", "Automated tool changer for the first build; production certification; detailed electrical safety design.", "proposed", "REF-README", "Manual tool change is the current project direction; confirm exclusions."],
    ["length_units", "mm", "draft", "", "Use mm for CAD geometry unless a field states otherwise."],
    ["mass_units", "kg", "draft", "", "Units for mass quantities."],
    ["force_units", "N", "draft", "", "Units for force quantities."],
    ["angle_units", "deg", "draft", "", "Units for rotations and angles."],
    ["time_units", "s", "draft", "", "Units for time quantities."],
    ["coordinate_convention", "Z up; define machine X/Y orientation before assembly modeling.", "draft", "", "Part-local datums should be explicit; do not infer global orientation."],
    ["operating_environment", "Workshop; plasma cutting creates substantial EMI, heat, dross, and water exposure around the bed.", "draft", "REF-README", "Detailed environmental limits remain open."],
    ["material_assumptions", "Tank plate is documented as 6 mm steel; most other materials are undecided.", "draft", "REF-README", "Do not infer steel grade or material properties."],
    ["safety_notes", "Lift/counterweight, electrical isolation, plasma EMI, pinch points, and table stability need engineering review.", "draft", "REF-README", "CAD geometry and spreadsheet checks are not safety certification."],
]

CONSTRAINT_HEADERS = [
    "constraint_id", "scope", "target_id", "name", "operator", "value",
    "minimum", "maximum", "units", "priority", "status", "rationale",
    "source_ref",
]
CONSTRAINTS = [
    ["CON-DEVICE-XMAX", "device", "", "Maximum outer X envelope", "<=", 1600, "", "", "mm", "hard", "confirmed", "Customer outer envelope target.", "REF-README"],
    ["CON-DEVICE-YMAX", "device", "", "Maximum outer Y envelope", "<=", 2500, "", "", "mm", "hard", "confirmed", "Customer outer envelope target.", "REF-README"],
    ["CON-SHEET-SHORTMIN", "device", "", "Minimum sheet short dimension", ">=", 1000, "", "", "mm", "hard", "confirmed", "Minimum sheet size requested by customer.", "REF-README"],
    ["CON-SHEET-LONGMIN", "device", "", "Minimum sheet long dimension", ">=", 2000, "", "", "mm", "hard", "confirmed", "Minimum sheet size requested by customer.", "REF-README"],
    ["CON-TANK-INTERNAL-W", "part", "PART-TANK", "Tank internal width", "=", 1712, "", "", "mm", "hard", "draft", "Documented square water-table interior; reconcile with rectangular machine envelope.", "REF-README"],
    ["CON-TANK-INTERNAL-L", "part", "PART-TANK", "Tank internal length", "=", 1712, "", "", "mm", "hard", "draft", "Documented tank is square; sheet support and machine envelope require confirmation.", "REF-README"],
    ["CON-TANK-DEPTH", "part", "PART-TANK", "Internal tank depth", "=", 140, "", "", "mm", "hard", "draft", "Current documented depth; water mass implications need updated lift analysis.", "REF-README"],
    ["CON-TANK-PLATE", "part", "PART-TANK", "Tank plate thickness", "=", 6, "", "", "mm", "hard", "draft", "Documented plate thickness; steel grade unspecified.", "REF-README"],
    ["CON-SLAT-THICKNESS", "part", "PART-SLAT", "Slat thickness", "=", 3, "", "", "mm", "preferred", "proposed", "README notes 2 mm is flimsy and suggests 3 mm if available.", "REF-README"],
    ["CON-SLAT-HEIGHT", "part", "PART-SLAT", "Slat depth", "=", 40, "", "", "mm", "hard", "draft", "Documented slat depth.", "REF-README"],
    ["CON-SLAT-SPACING", "part", "PART-SLAT", "Maximum slat spacing", "<=", 50, "", "", "mm", "hard", "draft", "Documented maximum spacing.", "REF-README"],
    ["CON-SLAT-TOP", "feature", "FEAT-SLAT-TOP", "Slat top below tank rim", "=", 25, "", "", "mm", "hard", "draft", "One groove height is intended to support the documented water-level modes.", "REF-README"],
    ["CON-XY-RAPID", "device", "", "Desired plasma travel speed", ">=", 15, "", "", "m/min", "preferred", "draft", "Sizing notes use 15 m/min; confirm whether this is a hard requirement.", "REF-AXIS-SIZING"],
    ["CON-Z-FEED", "device", "", "Z drilling design feed force", ">=", 2000, "", "", "N", "preferred", "draft", "Design value applies to Z drilling, not X/Y axes.", "REF-Z-CALC"],
    ["CON-HG20", "device", "", "Linear guide family", "=", "HG20", "", "", "", "preferred", "draft", "Current X/Y guideway direction; final rail/block layout is open.", "REF-AXIS-SIZING"],
    ["CON-CTRL", "device", "", "Motion/control stack", "=", "LinuxCNC Qtplasmac + Mesa 7i76e + THCAD-5/10", "", "", "", "preferred", "draft", "Architecture captured in notes; verify exact Mesa hardware inventory.", "REF-CONTROL"],
]

PART_HEADERS = [
    "part_id", "name", "description", "role", "status", "priority",
    "model_type", "primary_datum", "target_envelope_x", "target_envelope_y",
    "target_envelope_z", "spec_notes", "reference_ids",
]
PARTS = [
    ["PART-TANK", "Water tank and pan", "Welded water-table pan with base and four side walls.", "Supports workpiece and contains cutting water.", "specified-draft", "1", "single-solid", "bottom-outside-center", "", "", "", "Internal dimensions currently documented as 1712 x 1712 x 140 mm; verify assembly clearances and weld design.", "REF-README"],
    ["PART-SLAT", "Replaceable slat grid", "Interlocking support fingers with toothed workpiece-contact edges.", "Supports sheet while allowing dross and water flow.", "specified-draft", "2", "multi-solid-part", "tank-bottom-center", "", "", "", "40 mm deep; 3 mm thickness is a preferred revision from the older 2 mm note; max spacing 50 mm. Slot direction and tooth dimensions need confirmation.", "REF-README"],
    ["PART-FRAME", "Table support frame and legs", "Frame carrying tank, gantry loads, and lift reactions.", "Supports table and provides structural interfaces.", "candidate", "3", "assembly", "base-frame-origin", "", "", "", "Member sections, material, bracing, leveling feet, and loads are not specified.", "REF-README"],
    ["PART-LIFT", "Table height lift", "Telescoping guides, cable/pulley assistance, counterweight, and positive locking pins.", "Raises/lowers table and holds selected height.", "candidate", "4", "assembly", "lift-base-datum", "", "", "", "Concept notes describe square-tube guides, cable/counterweight, and pin locking. Prior lift calculations assumed a shallower 100 mm water depth and need updating.", "REF-README"],
    ["PART-GANTRY", "XY gantry and motion system", "Long-axis rails/drive, cross-axis gantry, carriage, and motors.", "Moves the tool across the work envelope.", "candidate", "5", "assembly", "machine-origin", "", "", "", "HG20 guides and closed-loop NEMA 24 motors are under evaluation. Long axis rack preferred for speed; short-axis drive decision remains open.", "REF-AXIS-SIZING"],
    ["PART-TORCH-Z", "Torch Z and breakaway mount", "Vertical torch slide, clamp/interface, and collision breakaway.", "Positions and protects the plasma torch.", "candidate", "6", "assembly", "torch-axis-datum", "", "", "", "Hypertherm torch has its own positioning sleeve/rack; clamp and Z arrangement need physical dimensions and confirmation.", "REF-TORCH-SPEC"],
    ["PART-TORCH", "Powermax 65 machine torch", "Customer-supplied Duramax 180-degree full-length machine torch.", "Plasma cutting tool.", "vendor-reference", "7", "vendor-component", "torch-mount-datum", "", "", "", "Reference component only until exact torch variant and mounting dimensions are confirmed.", "REF-TORCH-SPEC"],
    ["PART-DRILL", "Manual-change drill spindle head", "Optional drill/spindle tool head with separate motor and belt-driven spindle concept.", "Future drilling and light machining tool.", "candidate", "8", "assembly", "spindle-axis-datum", "", "", "", "2.5 mm Z screw and Flipsky spindle motor are concepts in notes; do not combine the plasma Z and drilling Z until architecture is decided.", "REF-Z-CALC;REF-FLIPSKY"],
    ["PART-CONTROL", "CNC control and electrical system", "LinuxCNC/Mesa motion control, THC sensing, wiring, and cabinet.", "Controls motion, torch interface, and height sensing.", "candidate", "9", "assembly", "cabinet-datum", "", "", "", "Electrical layout and exact Mesa boards are not verified. Treat as a non-geometric assembly until detailed requirements exist.", "REF-CONTROL"],
    ["PART-ROTARY", "Tube rotary axis", "Optional fourth-axis fixture for tube plasma cutting.", "Supports and rotates tube stock.", "future", "10", "assembly", "rotary-axis-datum", "", "", "", "Customer wishlist only; tube range, chuck/support, drive, and removable mounting interface are undefined.", "REF-README"],
]

FEATURE_HEADERS = [
    "feature_id", "part_id", "name", "description", "required",
    "feature_kind", "dimensions_or_intent", "interface_name",
    "constraint_ids", "status", "notes",
]
FEATURES = [
    ["FEAT-TANK-BASE", "PART-TANK", "Tank base plate", "Flat 6 mm plate forming the bottom of the water pan.", "yes", "plate", "1712 x 1712 mm internal dimensions; outside size and bend/weld allowance need definition.", "", "CON-TANK-INTERNAL-W;CON-TANK-INTERNAL-L;CON-TANK-PLATE", "draft", "Geometry must distinguish inside dimensions from outside dimensions."],
    ["FEAT-TANK-WALLS", "PART-TANK", "Tank side walls", "Four walls enclosing the water pan.", "yes", "wall", "140 mm internal depth; wall joints and corner details unresolved.", "", "CON-TANK-DEPTH;CON-TANK-PLATE", "draft", "Do not infer weld prep or watertight construction details."],
    ["FEAT-TANK-DRAIN", "PART-TANK", "Drain outlet", "Low-point outlet for draining water and dross.", "yes", "port", "At lowest floor-pan point; port size and valve interface TBD.", "DrainPort", "", "open", "Exact outlet orientation must be clarified before detailed CAD."],
    ["FEAT-SLAT-TOOTH", "PART-SLAT", "Shark-tooth contact edge", "Narrow peaks provide discrete sheet contact.", "yes", "profile", "Tooth pitch, height, tip width, and manufacturable profile TBD.", "WorkpieceSupport", "CON-SLAT-THICKNESS;CON-SLAT-HEIGHT", "draft", "README describes shark-tooth style; no dimensioned tooth profile."],
    ["FEAT-SLAT-SLOTS", "PART-SLAT", "Interlocking slots", "Crossing slots allow the two slat directions to interlock.", "yes", "slot", "Slot depth and fit clearance TBD; one slot set is from bottom, the crossing set from top.", "SlatIntersection", "CON-SLAT-THICKNESS;CON-SLAT-HEIGHT", "draft", "Slot pattern and slat count require definition."],
    ["FEAT-SLAT-TOP", "PART-SLAT", "Slat elevation", "Sets the slat peaks below the tank rim.", "yes", "datum", "Top contact peak 25 mm below tank rim.", "SlatTopPlane", "CON-SLAT-TOP", "draft", "Height is from tooth peaks, not slat root."],
    ["FEAT-FRAME-MOUNTS", "PART-FRAME", "Tank mounting interfaces", "Interfaces carrying tank and lift loads.", "yes", "mounting-interface", "Locations, bolt pattern, and support method TBD.", "TankMount", "", "open", "Must coordinate with tank shell and lift."],
    ["FEAT-FRAME-GANTRY", "PART-FRAME", "Gantry rail interfaces", "Reference surfaces for long-axis guide rails.", "yes", "mounting-interface", "Rail spacing, straightness, and mounting arrangement TBD.", "LongAxisRailMount", "", "open", "HG20 is a candidate series, not a released rail layout."],
    ["FEAT-LIFT-GUIDES", "PART-LIFT", "Telescoping guide columns", "Guide table motion while maintaining alignment.", "yes", "linear-guide", "Nested square tubes with bearing guidance; sizes and stroke TBD.", "LiftGuideAxis", "", "draft", "Pinch points and buckling/stability require engineering review."],
    ["FEAT-LIFT-CABLE", "PART-LIFT", "Cable and pulley assist", "Counterweight-assisted manual lift concept.", "yes", "lift-interface", "Cable routing, pulley geometry, and rated hardware TBD.", "LiftCablePath", "", "draft", "Existing lift calculations predate current 140 mm tank depth."],
    ["FEAT-LIFT-PINS", "PART-LIFT", "Height lock pins", "Positive mechanical locks at discrete table heights.", "yes", "locking-interface", "Pin count, size, material, and hole positions TBD.", "HeightLock", "", "draft", "Do not treat counterweight/cable as the sole support at rest."],
    ["FEAT-GANTRY-RAILS", "PART-GANTRY", "HG20 guide rails", "Linear guidance along machine axes.", "yes", "linear-guide", "Two rails and two blocks per axis are documented concept; lengths and spacing TBD.", "RailDatum", "CON-HG20", "draft", "Long axis travel/speed and rail alignment require sizing."],
    ["FEAT-GANTRY-DRIVE", "PART-GANTRY", "Rack or ballscrew drives", "Transmits motor torque to axis carriages.", "yes", "linear-drive", "Rack preferred for long axis; short axis drive not decided.", "DriveAxis", "CON-XY-RAPID", "draft", "Do not assume a final tooth count or screw lead."],
    ["FEAT-GANTRY-MOTOR", "PART-GANTRY", "Closed-loop stepper motors", "Axis actuation with encoder feedback.", "yes", "motor-mount", "NEMA 24, 4 Nm motor under evaluation.", "MotorMount", "", "draft", "Exact motor SKU/mounting pattern should be referenced."],
    ["FEAT-TORCHZ-SLIDE", "PART-TORCH-Z", "Torch vertical slide", "Provides torch height positioning and THC movement.", "yes", "linear-slide", "Stroke and guide arrangement TBD.", "TorchZAxis", "", "open", "Plasma Z architecture is not the drilling feed axis by assumption."],
    ["FEAT-TORCHZ-BREAKAWAY", "PART-TORCH-Z", "Breakaway interface", "Detects or limits damage from torch collision.", "yes", "safety-interface", "Trip direction, repeatability, sensing, and reset method TBD.", "BreakawayMount", "", "draft", "Physical mechanism is not dimensioned."],
    ["FEAT-TORCHZ-CLAMP", "PART-TORCH-Z", "Torch clamp", "Holds the machine torch sleeve in alignment.", "yes", "clamp", "Barrel and sleeve diameters must come from verified torch variant.", "TorchMount", "", "open", "Use authoritative torch drawing or measurement."],
    ["FEAT-TORCH-SLEEVE-RACK", "PART-TORCH", "Positioning sleeve rack", "Vendor rack interface associated with torch Z positioning.", "yes", "vendor-interface", "32 pitch/module approximately 0.8 mm per project notes; verify exact torch part.", "TorchRack", "", "draft", "Reference spec is project research, not verified customer hardware."],
    ["FEAT-DRILL-SPINDLE", "PART-DRILL", "Drill spindle", "Optional spindle for small-hole drilling and light machining.", "no", "spindle", "6 mm maximum tool; target about 1300 rpm from sizing notes.", "SpindleAxis", "", "draft", "Tool, bearing, and interface architecture unresolved."],
    ["FEAT-DRILL-Z", "PART-DRILL", "Drill feed axis", "Optional high-force Z feed for drilling.", "no", "linear-drive", "2.5 mm self-locking lead screw concept; design feed 2000 N.", "DrillZAxis", "CON-Z-FEED", "draft", "Do not merge with plasma torch Z without explicit architecture decision."],
    ["FEAT-CONTROL-THC", "PART-CONTROL", "Torch height sensing", "Arc voltage and ohmic sensing front-end.", "yes", "electrical-interface", "Mesa THCAD-10 arc voltage and THCAD-5 touch-off are documented selection.", "THCInterface", "CON-CTRL", "draft", "High-voltage interface must follow manufacturer documentation; CAD model is schematic only."],
    ["FEAT-CONTROL-MOTION", "PART-CONTROL", "Motion controller", "LinuxCNC/Mesa control stack for coordinated motion.", "yes", "electrical-interface", "LinuxCNC Qtplasmac and Mesa 7i76e documented direction.", "MotionIO", "CON-CTRL", "draft", "Verify exact board inventory and I/O mapping."],
    ["FEAT-ROTARY-CHUCK", "PART-ROTARY", "Tube support and rotation", "Supports and rotates tube stock for fourth-axis cutting.", "no", "rotary-interface", "Stock diameter, length, chuck, tail support, and axis position TBD.", "TubeAxis", "", "open", "Wishlist only."],
]

ASSEMBLY_HEADERS = [
    "assembly_id", "name", "description", "status", "priority", "notes",
]
ASSEMBLIES = [
    ["ASMB-TABLE", "Water table and lift", "Tank, slats, frame, and adjustable-height lift.", "candidate", "1", "Interfaces and lifting sequence are not yet defined."],
    ["ASMB-MOTION", "XY gantry and motion", "Machine frame, guides, drives, motors, and moving gantry.", "candidate", "2", "Axis orientation and mounting datums must be resolved."],
    ["ASMB-PLASMA", "Plasma tool head", "Torch Z/breakaway assembly carrying the vendor torch.", "candidate", "3", "Torch variant, stroke, and clamp dimensions need confirmation."],
    ["ASMB-CONTROL", "Control and electrical system", "Motion control, THC, torch signals, and electrical cabinet.", "candidate", "4", "Non-geometric connections and electrical interfaces need separate definitions."],
    ["ASMB-DRILL", "Optional drilling tool head", "Manual-change spindle and drill-feed concept.", "future", "5", "Plasma and drill tool-axis arrangement is undecided."],
    ["ASMB-ROTARY", "Optional tube axis", "Fourth-axis tube support and rotation assembly.", "future", "6", "Mounting position and envelope are undefined."],
]

INSTANCE_HEADERS = [
    "instance_id", "assembly_id", "part_id", "name", "x", "y", "z",
    "rotation_x", "rotation_y", "rotation_z", "placement_status", "notes",
]
INSTANCES = [
    ["INST-TABLE-TANK", "ASMB-TABLE", "PART-TANK", "Tank", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-TABLE-SLATS", "ASMB-TABLE", "PART-SLAT", "Slat grid", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-TABLE-FRAME", "ASMB-TABLE", "PART-FRAME", "Support frame", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-TABLE-LIFT", "ASMB-TABLE", "PART-LIFT", "Height lift", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-MOTION-FRAME", "ASMB-MOTION", "PART-FRAME", "Machine frame", "", "", "", "", "", "", "undefined", "Candidate membership only; frame scope may need splitting."],
    ["INST-MOTION-GANTRY", "ASMB-MOTION", "PART-GANTRY", "XY gantry", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-PLASMA-Z", "ASMB-PLASMA", "PART-TORCH-Z", "Torch Z assembly", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-PLASMA-TORCH", "ASMB-PLASMA", "PART-TORCH", "Machine torch", "", "", "", "", "", "", "undefined", "Candidate membership only; placement not designed."],
    ["INST-CONTROL", "ASMB-CONTROL", "PART-CONTROL", "Control system", "", "", "", "", "", "", "undefined", "Candidate membership only; cabinet placement not designed."],
    ["INST-DRILL-HEAD", "ASMB-DRILL", "PART-DRILL", "Drill head", "", "", "", "", "", "", "undefined", "Future candidate; not approved for initial build."],
    ["INST-DRILL-GANTRY", "ASMB-DRILL", "PART-GANTRY", "Supporting gantry", "", "", "", "", "", "", "undefined", "Future candidate; final tool interface is open."],
    ["INST-ROTARY", "ASMB-ROTARY", "PART-ROTARY", "Tube rotary", "", "", "", "", "", "", "undefined", "Future candidate; no placement or interface defined."],
    ["INST-ROTARY-GANTRY", "ASMB-ROTARY", "PART-GANTRY", "Supporting gantry", "", "", "", "", "", "", "undefined", "Future candidate; placement not designed."],
]

CONNECTION_HEADERS = [
    "connection_id", "assembly_id", "parent_instance_id", "parent_interface",
    "child_instance_id", "child_interface", "joint_type", "axis_or_datum",
    "lower_limit", "upper_limit", "units", "status", "notes",
]
CONNECTIONS: list[list[object]] = []

REFERENCE_HEADERS = [
    "reference_id", "part_id", "path", "document_type", "description",
    "authority", "notes",
]
REFERENCES = [
    ["REF-README", "", "README.md", "project-summary", "Project brief, current envelope, water-table dimensions, and decision log.", "project-notes", "Contains a mixture of confirmed customer data and evolving design notes."],
    ["REF-AXIS-SIZING", "PART-GANTRY", "calc/axis-sizing.md", "engineering-calculation", "X/Y sizing, drive comparison, and speed considerations.", "engineering-note", "Review assumptions before treating calculations as released design."],
    ["REF-Z-CALC", "PART-DRILL", "calc/z-axis-spindle.md", "engineering-calculation", "Drilling Z and optional spindle concept.", "engineering-note", "Concept calculations only."],
    ["REF-CONTROL", "PART-CONTROL", "calc/control-architecture.md", "architecture-note", "LinuxCNC, Mesa, and THC architecture.", "engineering-note", "Verify hardware inventory and electrical implementation."],
    ["REF-TORCH-SPEC", "PART-TORCH", "drawings/powermax65_cnc_torch_spec.md", "vendor-reference-summary", "Torch configuration and CNC interfaces.", "research-summary", "Verify exact customer torch model against primary vendor documentation."],
    ["REF-FLIPSKY", "PART-DRILL", "drawings/flipsky_7070_140kv_spec.md", "vendor-reference-summary", "Optional BLDC spindle motor dimensions and specifications.", "research-summary", "Motor suitability and actual variant need confirmation."],
    ["REF-BED-CAD", "PART-TANK", "drawings/PlasmaTable.FCStd", "freecad-document", "Existing FreeCAD document containing early tank/slat assembly geometry.", "project-model", "Open and inspect before reusing; it is not guaranteed to match this workbook exactly."],
    ["REF-MOTOR-STEP", "PART-GANTRY", "drawings/24HS40-5004D-E1000.STEP", "vendor-cad", "NEMA 24 motor STEP model.", "vendor-model", "Check orientation and exact motor variant."],
]

CHANGE_HEADERS = [
    "change_id", "date", "actor", "related_ids", "action", "rationale",
    "status",
]
CHANGES: list[list[object]] = []

SHEETS = [
    ("Device_Brief", ["field", "value", "status", "source_ref", "notes"], DEVICE_BRIEF),
    ("Constraints", CONSTRAINT_HEADERS, CONSTRAINTS),
    ("Parts", PART_HEADERS, PARTS),
    ("Features", FEATURE_HEADERS, FEATURES),
    ("Assemblies", ASSEMBLY_HEADERS, ASSEMBLIES),
    ("Instances", INSTANCE_HEADERS, INSTANCES),
    ("Connections", CONNECTION_HEADERS, CONNECTIONS),
    ("References", REFERENCE_HEADERS, REFERENCES),
    ("Change_Log", CHANGE_HEADERS, CHANGES),
]


def add_start_here(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Start_Here", 0)
    rows = [
        ["AI_CAD workflow workbook — draft"],
        ["Purpose", "Enter machine intent and constraints here. The exporter turns the data tables into JSON for the AI-assisted FreeCAD workflow."],
        ["Workflow", "1. Complete Device_Brief. 2. Add one row per requirement in Constraints. 3. Register Parts and Features. 4. Describe candidate groupings in Assemblies and occurrences in Instances. 5. Add named-interface joints in Connections when known. 6. Register supporting files in References. 7. Ask the agent to validate/export and work on a part."],
        ["Important", "This workbook contains a projected plasma-table draft. Verify status and source for every important value. Candidate/future rows are not approved design decisions."],
        ["IDs", "Use stable unique IDs: PART-, FEAT-, CON-, ASMB-, INST-, CONN-, REF-, CHG-. Do not reuse IDs after deleting a row."],
        ["Unknown values", "Leave unknown cells blank. Blank is not zero. Ask for clarification when a missing value changes geometry, fit, envelope, or safety."],
        ["Status", "draft/proposed/open/future are unresolved. confirmed means the user accepted the design intent, not that it is engineered or certified."],
        ["Geometry", "Keep part-local geometry separate from assembly placements. Use named interfaces and datums rather than relying on anonymous faces."],
        ["References", "Use repository-relative paths and identify document_type and authority. The agent must inspect the referenced file; a path alone is not evidence."],
        ["Change log", "When the agent adds or changes rows, it should record who/what changed, affected IDs, and rationale in Change_Log."],
        ["Export", "From the repository root run: python AI_CAD\\tools\\export_workbook.py. Generated JSON is derived output; do not edit it manually."],
        ["Limitations", "The exporter validates workbook structure and references. It does not prove dimensions, strength, safety, manufacturability, or FreeCAD model correctness."],
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 112
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.merge_cells("A1:B1")
    for row_index in range(2, len(rows) + 1):
        sheet.cell(row_index, 1).font = Font(bold=True, color="1F1F1F")
        sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        sheet.cell(row_index, 1).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(row_index, 2).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row_index].height = 42


def add_table_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    if rows:
        for row in rows:
            sheet.append(row)
    else:
        sheet.append([None] * len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 30
    for column_cells in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column_cells]
        width = min(max(max(values, default=10) + 2, 12), 46)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(displayName="tbl" + title.replace("_", ""), ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def add_validations(workbook: Workbook) -> None:
    validation_specs = {
        "Device_Brief": {"C": '"draft,proposed,confirmed,open"'},
        "Constraints": {
            "B": '"device,part,feature"',
            "E": '"=,<=,>=,range,one-of"',
            "J": '"hard,preferred,informational"',
            "K": '"draft,proposed,confirmed,open,waived"',
        },
        "Parts": {
            "E": '"candidate,specified-draft,vendor-reference,confirmed,modelled,reviewed,deferred,future"',
        },
        "Features": {
            "E": '"yes,no,undecided"',
            "J": '"draft,proposed,confirmed,open,modelled,reviewed,deferred"',
        },
        "Assemblies": {"D": '"candidate,draft,confirmed,modelled,reviewed,deferred,future"'},
        "Instances": {"K": '"undefined,proposed,placed,verified"'},
        "Connections": {
            "G": '"fixed,slider,revolute,cylindrical,planar,gear,other"',
            "L": '"draft,proposed,confirmed,verified,open,waived"',
        },
        "References": {"F": '"manufacturer,customer,project-model,project-notes,engineering-note,research-summary,user-sketch,approximate"'},
        "Change_Log": {"G": '"proposed,accepted,rejected,implemented"'},
    }
    for sheet_name, columns in validation_specs.items():
        sheet = workbook[sheet_name]
        for column, formula in columns.items():
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.error = "Choose a listed value or leave blank."
            validation.errorTitle = "Value not in list"
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}1000")


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_start_here(workbook)
    for title, headers, rows in SHEETS:
        add_table_sheet(workbook, title, headers, rows)
    add_validations(workbook)
    workbook.properties.title = "AI_CAD plasma table workflow draft"
    workbook.properties.subject = "Human-edited machine requirements and part registry"
    workbook.properties.creator = "AI_CAD workflow"
    workbook.save(WORKBOOK_PATH)
    print(f"Created {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
