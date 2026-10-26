"""Validate the AI_CAD workbook and export its tables as JSON."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = REPOSITORY_ROOT / "AI_CAD" / "workflow.xlsx"
DEFAULT_OUTPUT = REPOSITORY_ROOT / "AI_CAD" / "generated"

REQUIRED_COLUMNS = {
    "Device_Brief": {"field", "value", "status", "source_ref", "notes"},
    "Constraints": {
        "constraint_id", "scope", "target_id", "name", "operator", "value",
        "minimum", "maximum", "units", "priority", "status", "rationale",
        "source_ref",
    },
    "Parts": {
        "part_id", "name", "description", "role", "status", "priority",
        "model_type", "primary_datum", "target_envelope_x",
        "target_envelope_y", "target_envelope_z", "spec_notes", "reference_ids",
    },
    "Features": {
        "feature_id", "part_id", "name", "description", "required",
        "feature_kind", "dimensions_or_intent", "interface_name",
        "constraint_ids", "status", "notes",
    },
    "Assemblies": {
        "assembly_id", "name", "description", "status", "priority", "notes",
    },
    "Instances": {
        "instance_id", "assembly_id", "part_id", "name", "x", "y", "z",
        "rotation_x", "rotation_y", "rotation_z", "placement_status", "notes",
    },
    "Connections": {
        "connection_id", "assembly_id", "parent_instance_id",
        "parent_interface", "child_instance_id", "child_interface",
        "joint_type", "axis_or_datum", "lower_limit", "upper_limit", "units",
        "status", "notes",
    },
    "References": {
        "reference_id", "part_id", "path", "document_type", "description",
        "authority", "notes",
    },
    "Change_Log": {
        "change_id", "date", "actor", "related_ids", "action", "rationale",
        "status",
    },
}
ID_COLUMNS = {
    "Constraints": "constraint_id",
    "Parts": "part_id",
    "Features": "feature_id",
    "Assemblies": "assembly_id",
    "Instances": "instance_id",
    "Connections": "connection_id",
    "References": "reference_id",
    "Change_Log": "change_id",
}


class WorkbookValidationError(ValueError):
    """Raised when workbook content cannot be safely exported."""


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def split_ids(value: Any) -> list[str]:
    if value is None or str(value).strip() == "":
        return []
    return [item.strip() for item in str(value).split(";") if item.strip()]


def read_table(workbook, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise WorkbookValidationError(f"Missing required worksheet: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    if any(not isinstance(header, str) or not header.strip() for header in headers):
        raise WorkbookValidationError(f"{sheet_name}: header row has blank/non-text cells")
    if len(set(headers)) != len(headers):
        raise WorkbookValidationError(f"{sheet_name}: duplicate column headers")
    missing = REQUIRED_COLUMNS[sheet_name] - set(headers)
    if missing:
        raise WorkbookValidationError(
            f"{sheet_name}: missing columns: {', '.join(sorted(missing))}"
        )

    records = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        record = {
            header: json_value(value)
            for header, value in zip(headers, values)
        }
        if all(value is None or str(value).strip() == "" for value in record.values()):
            continue
        record["_row"] = row_number
        records.append(record)
    return records


def validate_unique_ids(tables: dict[str, list[dict[str, Any]]]) -> None:
    for sheet_name, id_column in ID_COLUMNS.items():
        seen: dict[str, int] = {}
        for record in tables[sheet_name]:
            raw_id = record.get(id_column)
            if raw_id is None or not str(raw_id).strip():
                raise WorkbookValidationError(
                    f"{sheet_name} row {record['_row']}: {id_column} is required"
                )
            identifier = str(raw_id).strip()
            if identifier in seen:
                raise WorkbookValidationError(
                    f"{sheet_name} row {record['_row']}: duplicate {identifier}; "
                    f"first used on row {seen[identifier]}"
                )
            seen[identifier] = record["_row"]


def validate_references(tables: dict[str, list[dict[str, Any]]]) -> None:
    part_ids = {str(row["part_id"]) for row in tables["Parts"]}
    feature_ids = {str(row["feature_id"]) for row in tables["Features"]}
    constraint_ids = {
        str(row["constraint_id"]) for row in tables["Constraints"]
    }
    assembly_ids = {str(row["assembly_id"]) for row in tables["Assemblies"]}
    reference_ids = {str(row["reference_id"]) for row in tables["References"]}
    instance_ids = {str(row["instance_id"]) for row in tables["Instances"]}
    instances_by_id = {
        str(row["instance_id"]): row for row in tables["Instances"]
    }

    for brief in tables["Device_Brief"]:
        source_ref = str(brief.get("source_ref") or "").strip()
        if source_ref and source_ref not in reference_ids:
            raise WorkbookValidationError(
                f"Device_Brief row {brief['_row']}: unknown source_ref "
                f"{source_ref!r}"
            )

    for feature in tables["Features"]:
        part_id = str(feature.get("part_id") or "").strip()
        if part_id not in part_ids:
            raise WorkbookValidationError(
                f"Features row {feature['_row']}: unknown part_id {part_id!r}"
            )
        for constraint_id in split_ids(feature.get("constraint_ids")):
            if constraint_id not in constraint_ids:
                raise WorkbookValidationError(
                    f"Features row {feature['_row']}: unknown constraint_id "
                    f"{constraint_id!r}"
                )
            linked_constraint = next(
                row for row in tables["Constraints"]
                if row["constraint_id"] == constraint_id
            )
            if (
                linked_constraint.get("scope") == "feature"
                and linked_constraint.get("target_id") != feature.get("feature_id")
            ):
                raise WorkbookValidationError(
                    f"Features row {feature['_row']}: constraint {constraint_id!r} "
                    "targets a different feature"
                )

    for constraint in tables["Constraints"]:
        scope = str(constraint.get("scope") or "").strip().lower()
        target_id = str(constraint.get("target_id") or "").strip()
        valid_targets = {"device": {""}, "part": part_ids, "feature": feature_ids}
        if scope not in valid_targets:
            raise WorkbookValidationError(
                f"Constraints row {constraint['_row']}: invalid scope {scope!r}"
            )
        if target_id not in valid_targets[scope]:
            raise WorkbookValidationError(
                f"Constraints row {constraint['_row']}: target {target_id!r} "
                f"does not exist for scope {scope!r}"
            )
        source_ref = str(constraint.get("source_ref") or "").strip()
        if source_ref and source_ref not in reference_ids:
            raise WorkbookValidationError(
                f"Constraints row {constraint['_row']}: unknown source_ref "
                f"{source_ref!r}"
            )

    for part in tables["Parts"]:
        for key in ("target_envelope_x", "target_envelope_y", "target_envelope_z"):
            value = part.get(key)
            if value not in (None, "") and (
                not isinstance(value, (int, float)) or value <= 0
            ):
                raise WorkbookValidationError(
                    f"Parts row {part['_row']}: {key} must be a positive number "
                    "in the project length unit"
                )
        for reference_id in split_ids(part.get("reference_ids")):
            if reference_id not in reference_ids:
                raise WorkbookValidationError(
                    f"Parts row {part['_row']}: unknown reference_id "
                    f"{reference_id!r}"
                )

    for reference in tables["References"]:
        part_id = str(reference.get("part_id") or "").strip()
        if part_id and part_id not in part_ids:
            raise WorkbookValidationError(
                f"References row {reference['_row']}: unknown part_id {part_id!r}"
            )
        relative_path = str(reference.get("path") or "").strip()
        if relative_path:
            resolved = (REPOSITORY_ROOT / relative_path).resolve()
            if not resolved.is_relative_to(REPOSITORY_ROOT.resolve()):
                raise WorkbookValidationError(
                    f"References row {reference['_row']}: path escapes repository: "
                    f"{relative_path!r}"
                )
            if not resolved.exists():
                raise WorkbookValidationError(
                    f"References row {reference['_row']}: file does not exist: "
                    f"{relative_path!r}"
                )

    for instance in tables["Instances"]:
        assembly_id = str(instance.get("assembly_id") or "").strip()
        part_id = str(instance.get("part_id") or "").strip()
        if assembly_id not in assembly_ids:
            raise WorkbookValidationError(
                f"Instances row {instance['_row']}: unknown assembly_id "
                f"{assembly_id!r}"
            )
        if part_id not in part_ids:
            raise WorkbookValidationError(
                f"Instances row {instance['_row']}: unknown part_id {part_id!r}"
            )
        if instance.get("placement_status") == "placed":
            position = [instance.get(axis) for axis in ("x", "y", "z")]
            rotation = [
                instance.get(axis)
                for axis in ("rotation_x", "rotation_y", "rotation_z")
            ]
            if any(not isinstance(value, (int, float)) for value in position + rotation):
                raise WorkbookValidationError(
                    f"Instances row {instance['_row']}: placed instances need "
                    "numeric x/y/z and rotation_x/y/z values"
                )

    for connection in tables["Connections"]:
        assembly_id = str(connection.get("assembly_id") or "").strip()
        parent_instance = str(connection.get("parent_instance_id") or "").strip()
        child_instance = str(connection.get("child_instance_id") or "").strip()
        if assembly_id not in assembly_ids:
            raise WorkbookValidationError(
                f"Connections row {connection['_row']}: unknown assembly_id "
                f"{assembly_id!r}"
            )
        if parent_instance not in instance_ids or child_instance not in instance_ids:
            raise WorkbookValidationError(
                f"Connections row {connection['_row']}: both instance IDs must "
                "exist in Instances"
            )
        if parent_instance == child_instance:
            raise WorkbookValidationError(
                f"Connections row {connection['_row']}: parent and child must differ"
            )
        if (
            instances_by_id[parent_instance].get("assembly_id") != assembly_id
            or instances_by_id[child_instance].get("assembly_id") != assembly_id
        ):
            raise WorkbookValidationError(
                f"Connections row {connection['_row']}: both instances must belong "
                f"to assembly {assembly_id!r}"
            )
        if connection.get("joint_type") not in {
            "fixed", "slider", "revolute", "cylindrical", "planar", "gear", "other"
        }:
            raise WorkbookValidationError(
                f"Connections row {connection['_row']}: unsupported joint_type"
            )

    for change in tables["Change_Log"]:
        raw_date = change.get("date")
        if raw_date is not None and not isinstance(raw_date, str):
            change["date"] = str(raw_date)


def clean_row(record: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in record.items() if key != "_row"}


def export(workbook_path: Path, output_dir: Path) -> None:
    if not workbook_path.is_file():
        raise WorkbookValidationError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        tables = {
            name: read_table(workbook, name)
            for name in REQUIRED_COLUMNS
        }
    finally:
        workbook.close()

    validate_unique_ids(tables)
    validate_references(tables)

    output_dir.mkdir(parents=True, exist_ok=True)
    brief = {
        str(row["field"]): {
            "value": row.get("value"),
            "status": row.get("status"),
            "source_ref": row.get("source_ref"),
            "notes": row.get("notes"),
        }
        for row in tables["Device_Brief"]
    }
    constraints = [clean_row(row) for row in tables["Constraints"]]
    references = [clean_row(row) for row in tables["References"]]
    parts = [clean_row(row) for row in tables["Parts"]]
    features = [clean_row(row) for row in tables["Features"]]
    assemblies = [clean_row(row) for row in tables["Assemblies"]]
    instances = [clean_row(row) for row in tables["Instances"]]
    connections = [clean_row(row) for row in tables["Connections"]]
    changes = [clean_row(row) for row in tables["Change_Log"]]
    parts_by_id = {str(part["part_id"]): part for part in parts}
    units = {
        "length": brief.get("length_units", {}).get("value") or "mm",
        "mass": brief.get("mass_units", {}).get("value") or "kg",
        "force": brief.get("force_units", {}).get("value") or "N",
        "angle": brief.get("angle_units", {}).get("value") or "deg",
        "time": brief.get("time_units", {}).get("value") or "s",
    }
    coordinate_convention = (
        brief.get("coordinate_convention", {}).get("value")
        or "Unspecified; clarify before assembly placement."
    )

    for part_id, part in parts_by_id.items():
        part_constraints = [
            row for row in constraints
            if row.get("scope") == "part" and row.get("target_id") == part_id
        ]
        part_features = [
            {
                **feature,
                "constraints": [
                    constraint
                    for constraint in constraints
                    if constraint["constraint_id"] in split_ids(
                        feature.get("constraint_ids")
                    )
                ],
            }
            for feature in features
            if feature.get("part_id") == part_id
        ]
        part_reference_ids = split_ids(part.get("reference_ids"))
        part_references = [
            reference for reference in references
            if reference["reference_id"] in part_reference_ids
            or reference.get("part_id") == part_id
        ]
        document = {
            "schemaVersion": 1,
            "units": units,
            "coordinateConvention": coordinate_convention,
            "targetEnvelope": {
                "x": part.get("target_envelope_x"),
                "y": part.get("target_envelope_y"),
                "z": part.get("target_envelope_z"),
                "units": units["length"],
                "status": "unspecified"
                if all(
                    part.get(key) in (None, "")
                    for key in (
                        "target_envelope_x",
                        "target_envelope_y",
                        "target_envelope_z",
                    )
                )
                else "user-target",
            },
            "part": part,
            "features": part_features,
            "constraints": part_constraints,
            "references": part_references,
            "provenance": {
                "workbook": workbook_path.relative_to(REPOSITORY_ROOT).as_posix(),
                "sheets": ["Parts", "Features", "Constraints", "References"],
            },
        }
        write_json(output_dir / "parts" / f"{part_id}.json", document)

    project_document = {
        "schemaVersion": 1,
        "units": units,
        "coordinateConvention": coordinate_convention,
        "device": brief,
        "constraints": constraints,
        "parts": parts,
        "features": features,
        "assemblies": assemblies,
        "instances": instances,
        "connections": connections,
        "references": references,
        "changeLog": changes,
        "provenance": {
            "workbook": workbook_path.relative_to(REPOSITORY_ROOT).as_posix(),
            "sheets": list(REQUIRED_COLUMNS),
        },
    }
    write_json(output_dir / "project.json", project_document)

    for assembly in assemblies:
        assembly_instances = [
            instance for instance in instances
            if instance.get("assembly_id") == assembly["assembly_id"]
        ]
        assembly_connections = [
            connection for connection in connections
            if connection.get("assembly_id") == assembly["assembly_id"]
        ]
        document = {
            "schemaVersion": 1,
            "assembly": assembly,
            "instances": [
                {
                    "instanceId": instance["instance_id"],
                    "partId": instance["part_id"],
                    "name": instance.get("name"),
                    "placement": {
                        "translation": {
                            "x": instance.get("x"),
                            "y": instance.get("y"),
                            "z": instance.get("z"),
                            "units": "mm",
                        },
                        "rotation": {
                            "x": instance.get("rotation_x"),
                            "y": instance.get("rotation_y"),
                            "z": instance.get("rotation_z"),
                            "units": "deg",
                        },
                    },
                    "placementStatus": instance.get("placement_status"),
                }
                for instance in assembly_instances
            ],
            "connections": assembly_connections,
            "provenance": {
                "workbook": workbook_path.relative_to(REPOSITORY_ROOT).as_posix(),
                "sheets": ["Assemblies", "Instances", "Connections"],
            },
        }
        write_json(output_dir / "assemblies" / f"{assembly['assembly_id']}.json", document)

    print(
        f"Exported {len(parts)} parts, {len(features)} features, "
        f"{len(constraints)} constraints, and {len(assemblies)} assemblies "
        f"to {output_dir}"
    )


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    try:
        export(args.workbook.resolve(), args.output.resolve())
    except (OSError, WorkbookValidationError) as error:
        print(f"AI_CAD export failed: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
