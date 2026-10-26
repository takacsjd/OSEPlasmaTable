from __future__ import annotations

import sys
import tempfile
import unittest
import json
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import export_workbook


class ExportWorkbookTests(unittest.TestCase):
    def test_seed_workbook_exports_expected_parts_and_assemblies(self) -> None:
        with tempfile.TemporaryDirectory(dir=export_workbook.REPOSITORY_ROOT) as temp:
            output = Path(temp) / "generated"
            export_workbook.export(export_workbook.DEFAULT_WORKBOOK, output)

            project = json.loads(
                (output / "project.json").read_text(encoding="utf-8")
            )
            self.assertGreaterEqual(len(project["parts"]), 8)
            self.assertGreaterEqual(len(project["features"]), 15)
            self.assertGreaterEqual(len(project["assemblies"]), 4)
            self.assertTrue((output / "parts" / "PART-TANK.json").is_file())
            self.assertTrue((output / "assemblies" / "ASMB-TABLE.json").is_file())
            tank = json.loads(
                (output / "parts" / "PART-TANK.json").read_text(encoding="utf-8")
            )
            self.assertEqual(tank["units"]["length"], "mm")
            self.assertEqual(tank["targetEnvelope"]["status"], "unspecified")

    def test_duplicate_part_id_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory(dir=export_workbook.REPOSITORY_ROOT) as temp:
            workbook_path = Path(temp) / "duplicate.xlsx"
            workbook = load_workbook(export_workbook.DEFAULT_WORKBOOK)
            sheet = workbook["Parts"]
            sheet.append([cell.value for cell in sheet[2]])
            workbook.save(workbook_path)
            workbook.close()

            workbook = load_workbook(workbook_path, data_only=True, read_only=True)
            try:
                tables = {
                    name: export_workbook.read_table(workbook, name)
                    for name in export_workbook.REQUIRED_COLUMNS
                }
            finally:
                workbook.close()
            with self.assertRaisesRegex(
                export_workbook.WorkbookValidationError, "duplicate PART-TANK"
            ):
                export_workbook.validate_unique_ids(tables)

    def test_workbook_contains_work_instructions_and_all_tables(self) -> None:
        workbook = load_workbook(export_workbook.DEFAULT_WORKBOOK, read_only=False)
        try:
            self.assertEqual(workbook.sheetnames[0], "Start_Here")
            self.assertEqual(
                set(export_workbook.REQUIRED_COLUMNS),
                set(workbook.sheetnames) - {"Start_Here"},
            )
            for sheet_name in export_workbook.REQUIRED_COLUMNS:
                self.assertTrue(workbook[sheet_name].tables)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
